"""Admin panelida Excel eksport imkoniyatini beruvchi universal mixin."""

import io

from django.contrib import admin, messages
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.urls import reverse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

_SESSION_IDS = "admin_excel_export_ids"
_SESSION_MODEL = "admin_excel_export_model"


def _get_concrete_fields(model):
    """Modelning DB ustuniga ega barcha fieldlarini qaytaradi."""
    result = []
    for f in model._meta.get_fields():
        if not hasattr(f, "column"):
            continue
        if getattr(f, "many_to_many", False):
            continue
        # ForeignKey uchun attname (masalan: role_id), oddiy field uchun name
        col_name = getattr(f, "attname", f.name)
        verbose = str(getattr(f, "verbose_name", col_name)).capitalize()
        result.append({"name": col_name, "verbose": verbose})
    return result


def _build_excel(model, ids, selected_fields):
    """Tanlangan fieldlar bo'yicha .xlsx fayl yaratadi va HttpResponse qaytaradi."""
    qs = model.objects.filter(pk__in=ids).values(*selected_fields)

    wb = Workbook()
    ws = wb.active
    ws.title = str(model._meta.verbose_name_plural)[:31]

    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")

    # Sarlavha qatori
    field_map = {f["name"]: f["verbose"] for f in _get_concrete_fields(model)}
    headers = [field_map.get(f, f) for f in selected_fields]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Ma'lumot qatorlari
    for row in qs:
        ws.append([("" if row[f] is None else str(row[f])) for f in selected_fields])

    # Ustun kengligi avtomatik
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"{model._meta.model_name}_export.xlsx"
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


class ExportExcelMixin:
    """Har qanday ModelAdmin ga Excel eksport qo'shuvchi mixin."""

    def get_urls(self):
        from django.urls import path

        meta = self.model._meta
        url_name = f"{meta.app_label}_{meta.model_name}_export_excel_fields"
        extra = [
            path(
                "export-excel-fields/",
                self.admin_site.admin_view(self._export_excel_fields_view),
                name=url_name,
            ),
        ]
        return extra + super().get_urls()

    @admin.action(description="Excel ga eksport qilish")
    def export_to_excel(self, request, queryset):
        """Tanlangan qatorlarni Excel ga eksport qilish uchun field tanlash sahifasiga yo'naltiradi."""
        ids = [str(pk) for pk in queryset.values_list("pk", flat=True)]
        if not ids:
            self.message_user(
                request, "Hech qanday yozuv tanlanmadi.", level=messages.WARNING
            )
            return

        request.session[_SESSION_IDS] = ids
        request.session[_SESSION_MODEL] = {
            "app_label": self.model._meta.app_label,
            "model_name": self.model._meta.model_name,
        }
        meta = self.model._meta
        url = reverse(f"admin:{meta.app_label}_{meta.model_name}_export_excel_fields")
        return redirect(url)

    def _export_excel_fields_view(self, request):
        """Field tanlash sahifasi: GET — forma, POST — Excel yuklab olish."""
        saved_model = request.session.get(_SESSION_MODEL, {})
        meta = self.model._meta

        # Sessiya boshqa model uchun ekanligini tekshirish
        if (
            saved_model.get("app_label") != meta.app_label
            or saved_model.get("model_name") != meta.model_name
        ):
            self.message_user(
                request,
                "Sessiya eskirgan yoki noto'g'ri model. Qaytadan tanlang.",
                level=messages.ERROR,
            )
            return redirect("../../")

        ids = request.session.get(_SESSION_IDS, [])
        fields = _get_concrete_fields(self.model)

        if request.method == "POST":
            selected = request.POST.getlist("fields")
            if not selected:
                self.message_user(
                    request, "Kamida 1 ta maydon tanlang.", level=messages.WARNING
                )
            else:
                return _build_excel(self.model, ids, selected)

        context = {
            **self.admin_site.each_context(request),
            "title": f"{meta.verbose_name_plural} — Excel eksport",
            "subtitle": f"{len(ids)} ta yozuv tanlangan",
            "fields": fields,
            "ids_count": len(ids),
            "opts": meta,
            "has_permission": True,
        }
        return render(request, "admin/export_excel_fields.html", context)
